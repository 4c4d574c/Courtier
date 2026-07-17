import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

DOMAIN_MAP = {
    "GOV": "GOV",
    "AUD": "AUD",
    "ECO": "ECON",
    "FIN": "FIN",
    "IND": "INDU",
    "COM": "TRADE",
    "EDU": "EDU",
    "SCI": "TECH",
    "LAB": "LIVELIHOOD",
    "HLT": "HEALTH",
    "RES": "RESOURCE",
    "ENV": "ECO_ENV",
    "WAT": "WATER",
    "TRA": "TRANSPORT",
    "URB": "URBAN",
    "AGR": "AGRI",
    "EMG": "EMERGENCY",
    "PUB": "SECURITY",
    "CUL": "CULTURE",
}

SEVERITY_MAP = {
    "禁止": "error",
    "不得": "error",
    "强制要求": "warning",
    "必须": "warning",
    "特别禁止": "error",
}

def determine_severity(text):
    if "特别禁止" in text:
        return "error"
    if "强制要求" in text or "必须" in text:
        return "warning"
    if "不得" in text or "禁止" in text:
        return "error"
    return "info"

def extract_rules(md_path):
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    rules = []
    current_domain = None

    lines = content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        domain_match = re.match(r"^## (\d+)\.\s+(\w+)\s+(.+?)领域", line)
        if domain_match:
            code = domain_match.group(2)
            current_domain = DOMAIN_MAP.get(code, code)
            i += 1
            continue

        if line.startswith("## 通用禁止规则"):
            current_domain = "GENERAL"
            i += 1
            continue

        if line.startswith("## 审核逻辑建议"):
            current_domain = None
            i += 1
            continue

        section_match = re.match(r"^### (.+)$", line)
        if section_match and current_domain:
            i += 1
            continue

        rule_match = re.match(r"^-\s+\*\*(.+?)\*\*(.*)$", line)
        if rule_match and current_domain:
            keyword = rule_match.group(1)
            rest = rule_match.group(2).strip()
            full_text = keyword + rest

            name = full_text
            if len(name) > 100:
                name = name[:97] + "..."

            severity = determine_severity(full_text)

            rules.append({
                "name": name,
                "domain_id": current_domain,
                "description": full_text,
                "severity": severity,
                "pattern": "",
            })

        i += 1

    return rules

def create_excel(rules, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "禁止规则清单"

    headers = ["name", "domain_id", "description", "severity", "pattern"]
    header_labels = ["规则名称", "领域ID", "规则描述", "严重程度", "匹配模式"]

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="微软雅黑", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap_align = Alignment(vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    severity_fills = {
        "error": PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid"),
        "warning": PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid"),
        "info": PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
    }

    for col_idx, label in enumerate(header_labels, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for row_idx, rule in enumerate(rules, 2):
        for col_idx, key in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=rule[key])
            cell.font = cell_font
            cell.border = thin_border
            if key in ("severity", "domain_id"):
                cell.alignment = center_align
            else:
                cell.alignment = wrap_align

            if key == "severity":
                fill = severity_fills.get(rule[key])
                if fill:
                    cell.fill = fill

    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 20

    wb.save(output_path)
    print(f"Excel saved to {output_path}")
    print(f"Total rules: {len(rules)}")

if __name__ == "__main__":
    md_path = "tests/assets/各领域禁止规则清单.md"
    output_path = "tests/assets/各领域禁止规则清单.xlsx"
    rules = extract_rules(md_path)
    create_excel(rules, output_path)
