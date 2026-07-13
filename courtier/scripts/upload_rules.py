import argparse
import sys
import time

import requests
from openpyxl import load_workbook

HEADERS = {"Content-Type": "application/json"}


def read_rules(xlsx_path: str) -> list[dict]:
    wb = load_workbook(xlsx_path, read_only=True)
    ws = wb.active

    headers_map = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(1, col).value
        if val:
            headers_map[col] = val

    col_keys = {}
    field_map = {
        "规则名称": "name",
        "领域ID": "domain_id",
        "规则描述": "description",
        "严重程度": "severity",
        "匹配模式": "pattern",
    }
    for col, label in headers_map.items():
        key = field_map.get(label)
        if key:
            col_keys[col] = key

    rules = []
    for row in range(2, ws.max_row + 1):
        rule = {}
        for col, key in col_keys.items():
            val = ws.cell(row, col).value
            rule[key] = val if val is not None else ""
        if rule.get("name") and rule.get("domain_id"):
            rules.append(rule)

    wb.close()
    return rules


def upload_rules(rules: list[dict], base_url: str, delay: float = 0.05):
    url = f"{base_url}/api/v1/rules/create"
    success = 0
    fail = 0
    skipped = 0

    for i, rule in enumerate(rules, 1):
        try:
            resp = requests.post(url, json=rule, headers=HEADERS, timeout=10)
            if resp.status_code == 200:
                success += 1
            else:
                fail += 1
                print(f"  [{i}] FAIL {resp.status_code}: {rule['name'][:50]} -> {resp.text[:100]}")
        except requests.RequestException as e:
            fail += 1
            print(f"  [{i}] ERROR: {rule['name'][:50]} -> {e}")

        if delay > 0:
            time.sleep(delay)

        if i % 20 == 0:
            print(f"  ... progress: {i}/{len(rules)} (ok={success}, fail={fail})")

    return success, fail


def main():
    parser = argparse.ArgumentParser(description="批量上传规则到 /api/v1/rules/create")
    parser.add_argument(
        "--xlsx",
        default="tests/assets/各领域禁止规则清单.xlsx",
        help="Excel 文件路径",
    )
    parser.add_argument(
        "--base-url",
        default="http://192.168.100.222:8000",
        help="API 服务地址",
    )
    parser.add_argument(
        "--delay",
        type=float,
        default=0.05,
        help="每次请求间隔秒数",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="仅解析 Excel 不实际上传",
    )
    args = parser.parse_args()

    print(f"Reading: {args.xlsx}")
    rules = read_rules(args.xlsx)
    print(f"Total rules parsed: {len(rules)}")

    if not rules:
        print("No rules found, exiting.")
        sys.exit(1)

    if args.dry_run:
        print("\n[DRY RUN] First 5 rules:")
        for r in rules[:5]:
            print(f"  [{r['domain_id']}] [{r['severity']}] {r['name'][:60]}")
        print(f"\n[DRY RUN] Would upload {len(rules)} rules to {args.base_url}")
        sys.exit(0)

    print(f"\nUploading to {args.base_url}/api/v1/rules/create ...")
    success, fail = upload_rules(rules, args.base_url, delay=args.delay)

    print(f"\nDone! success={success}, fail={fail}, total={len(rules)}")
    if fail > 0:
        sys.exit(1)


if __name__ == "__main__":
    main()
