"""Unit tests for resource library text extraction and chunking."""

from __future__ import annotations

import pytest

from courtier.agent.api.services.resource_service import (
    _extract_text,
    _split_chunks,
    _tail_paragraphs,
    build_chunk_actions,
)


class TestSplitChunks:
    def test_short_text_single_chunk(self):
        assert _split_chunks("第一段。\n第二段。") == ["第一段。\n第二段。"]

    def test_paragraphs_grouped_up_to_size(self):
        paras = [f"第{i}段" + "字" * 90 for i in range(10)]  # ~94 chars each
        chunks = _split_chunks("\n".join(paras), chunk_size=300)
        assert all(len(c) <= 300 for c in chunks)
        assert len(chunks) > 1
        # No paragraph content lost
        assert "".join(chunks).count("第") == 10

    def test_oversized_paragraph_hard_split(self):
        long_para = "长" * 2500
        chunks = _split_chunks(long_para, chunk_size=1000)
        assert chunks == ["长" * 1000, "长" * 1000, "长" * 500]

    def test_empty_and_blank_paragraphs_ignored(self):
        assert _split_chunks("\n\n  \n\n") == []

    def test_mixed_short_and_long_paragraphs(self):
        text = "短段。\n" + "长" * 1200 + "\n结尾。"
        chunks = _split_chunks(text, chunk_size=500)
        assert chunks[0] == "短段。"
        assert chunks[1] == "长" * 500
        assert chunks[2] == "长" * 500
        assert chunks[3] == "长" * 200 + "\n结尾。" or chunks[-1].endswith("结尾。")

    def test_overlap_prepends_previous_tail(self):
        paras = [f"第{i}段" + "字" * 90 for i in range(8)]
        chunks = _split_chunks("\n".join(paras), chunk_size=300, overlap=100)
        # Overlap duplicates content, so paragraph markers exceed the count.
        assert "".join(chunks).count("第") > 8
        for prev, nxt in zip(chunks, chunks[1:]):
            assert nxt.startswith(_tail_paragraphs(prev, 100))

    def test_overlap_off_by_default(self):
        paras = [f"第{i}段" + "字" * 90 for i in range(8)]
        assert _split_chunks("\n".join(paras), chunk_size=300) == _split_chunks(
            "\n".join(paras), chunk_size=300, overlap=0
        )


class TestTailParagraphs:
    def test_paragraph_aligned_tail(self):
        assert _tail_paragraphs("甲\n乙\n丙", 2) == "乙\n丙"
        assert _tail_paragraphs("甲\n乙\n丙", 3) == "甲\n乙\n丙"

    def test_single_oversized_paragraph_hard_cut(self):
        assert _tail_paragraphs("长" * 100, 10) == "长" * 10


class TestBuildChunkActions:
    def test_paragraph_index_matches_chunk_no(self):
        actions = build_chunk_actions(
            1,
            ["块一", "块二"],
            doc_type="txt",
            title="标题",
            author="",
            user_id="u",
            visibility="public",
            owner_id=None,
            tags=[],
            publish_date=None,
            index_name="idx",
        )
        bodies = [a for a in actions if "chunk_text" in a]
        assert bodies[0]["paragraph_index"] == 0
        assert bodies[1]["paragraph_index"] == 1
        assert bodies[0]["chunk_no"] == 0

    def test_vectors_attached_when_available(self):
        actions = build_chunk_actions(
            1,
            ["块一", "块二"],
            doc_type="txt",
            title="标题",
            author="",
            user_id="u",
            visibility="public",
            owner_id=None,
            tags=[],
            publish_date=None,
            index_name="idx",
            vectors=[[0.1, 0.2], None],
        )
        bodies = [a for a in actions if "chunk_text" in a]
        assert bodies[0]["chunk_vector"] == [0.1, 0.2]
        assert "chunk_vector" not in bodies[1]  # None entry stays lexical


class TestExtractText:
    def test_txt_file(self, tmp_path):
        p = tmp_path / "a.txt"
        p.write_text("纯文本内容\n第二行", encoding="utf-8")
        assert _extract_text(p) == "纯文本内容\n第二行"

    def test_md_file(self, tmp_path):
        p = tmp_path / "a.md"
        p.write_text("# 标题\n正文", encoding="utf-8")
        assert "标题" in _extract_text(p)

    def test_docx_file(self, tmp_path):
        import docx

        p = tmp_path / "a.docx"
        doc = docx.Document()
        doc.add_paragraph("第一段内容")
        doc.add_paragraph("第二段内容")
        doc.save(str(p))
        text = _extract_text(p)
        assert "第一段内容" in text
        assert "第二段内容" in text

    def test_pdf_file(self, tmp_path):
        import fitz

        p = tmp_path / "a.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "PDF 测试文本")
        doc.save(str(p))
        doc.close()
        assert "PDF" in _extract_text(p)


class TestUploadMagicByteGate:
    """Binary resource uploads must survive the magic-byte gate before
    reaching the parsers (PyMuPDF / python-docx / LibreOffice)."""

    class _FakeFile:
        def __init__(self, filename: str, payload: bytes):
            self.filename = filename
            self.payload = payload

        async def read(self, n: int = -1) -> bytes:
            chunk, self.payload = self.payload[:n], self.payload[n:]
            return chunk

    def _ingest(self, filename: str, payload: bytes):
        import asyncio
        from types import SimpleNamespace

        from fastapi import HTTPException

        from courtier.agent.api.services.resource_service import ingest_resource

        settings = SimpleNamespace(
            es_hosts="es:9200",
            minio_endpoint="",
            es_index_chunks="chunks",
            upload_dir="/tmp",
        )
        with pytest.raises(HTTPException) as exc_info:
            asyncio.run(
                ingest_resource(
                    self._FakeFile(filename, payload),
                    title="t",
                    author="",
                    source="",
                    tags="",
                    publish_date=None,
                    owner_name="u",
                    owner_id=1,
                    is_admin=False,
                    settings=settings,
                    db=None,
                )
            )
        return exc_info.value

    def test_fake_pdf_rejected(self):
        assert self._ingest("evil.pdf", b"not-a-pdf-at-all").status_code == 400

    def test_fake_docx_rejected(self):
        assert self._ingest("evil.docx", b"\x00" * 64).status_code == 400

    def test_magic_byte_matcher(self):
        from courtier.agent.api.services.file_service import (
            _content_matches_extension,
        )

        assert _content_matches_extension(".pdf", b"%PDF-1.7 rest")
        assert not _content_matches_extension(".pdf", b"junk-bytes")
        assert _content_matches_extension(".docx", b"PK\x03\x04rest")
        assert not _content_matches_extension(".docx", b"junk-bytes")
