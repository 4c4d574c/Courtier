"""validator 格式检查端到端测试（真实解析链路）。

链路：程序化生成 PDF（PyMuPDF/fitz）/ DOCX（python-docx）→ docparse.parse
→ ``Document.model_dump(exclude_none=True)``（与 parse_document 插件输出一致）
→ ``validator.validator(doc_dict, doc_type="通知")``。

覆盖：
- 合规 PDF（2 页、两页均带页码）：零间距误报、零页边距误报、零页码误报；
- 合规 DOCX：除图形要素（红色分隔线/印章/黑色反线——非文本行，规则引擎
  无法分类，缺失报错属预期）外零错误；
- 违规注入：(a) DOCX 正文段前距 20pt → 检出；(b) PDF 首页页码缺失、
  第 2 页页码字号错误 → 按修后语义检出（should_exist_on_page 只控制
  "缺失是否报错"，已存在的页码段在任何页都比较字体/字号）。

扫描件路径（OCR + LLM 结构识别）需要外部服务，无法在单元测试中走真实
链路，不做 e2e 覆盖。

测试环境限制：PDF 侧无「仿宋」字体文件可内嵌，PyMuPDF 只能使用 CJK 回退
字体（本机解析为 Heiti），故合规 PDF 的字体类报错不在断言范围内（DOCX
用例的字体名直接来自 run 属性，不受此限，做全量字体断言）。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import fitz
import pytest
from docparse import parse
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt

from validator.format_checker import validator

DOC_TYPE = "通知"
LINE_PITCH_PT = 28.95  # GB/T 9704 实践：全文固定行距 28.95pt

_A4_W, _A4_H = 595.28, 841.89
_LEFT_PT = 79.37  # 左边距 28mm
_TOP_PT = 105.0  # 上边距 ≈37mm

# 生成文档所用文字（避开规则引擎的标题/版记关键词误判）
_LOGO = "测试县人民政府文件"
_NUMBER = "测政发〔2026〕1号"
_TITLE = "关于开展测试工作的通知"
_ADDRESSEE = "各乡镇人民政府："
_SIGNATURE = "测试县人民政府"
_ISSUE_DATE = "2026年5月11日"
_OFFICE_LINE = "测试县人民政府办公室2026年5月12日印发"


def _audit(path: Path) -> tuple[Any, dict[str, Any]]:
    """走真实链路：docparse.parse → dict 化 → validator。"""
    document = parse(str(path))
    result = validator(document.model_dump(exclude_none=True), doc_type=DOC_TYPE)
    return document, result


def _errors(result: dict[str, Any], error_type: str) -> list[dict[str, Any]]:
    return [e for e in result["errors"] if e["error_type"] == error_type]


# ---------------------------------------------------------------------------
# PDF 生成
# ---------------------------------------------------------------------------


def _probe_line_metrics() -> dict[int, tuple[float, float]]:
    """实测各字号行 bbox：{size: (baseline→y0 距离, 行高)}。

    PyMuPDF 行高随运行环境的回退字体而定，用探针页实测后按目标 y0 排版，
    保证段间距测量值是确定性的。
    """
    probes = [(22, 120.0), (16, 220.0), (14, 320.0)]
    doc = fitz.open()
    page = doc.new_page(width=_A4_W, height=_A4_H)
    for size, baseline in probes:
        page.insert_text((_LEFT_PT, baseline), "国", fontname="china-s", fontsize=size)
    base_by_size = {size: base for size, base in probes}
    metrics: dict[int, tuple[float, float]] = {}
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block["lines"]:
            for span in line["spans"]:
                size = int(round(span["size"]))
                y0, y1 = line["bbox"][1], line["bbox"][3]
                metrics[size] = (base_by_size[size] - y0, y1 - y0)
    doc.close()
    assert len(metrics) == 3, f"探针行实测失败：{metrics}"
    return metrics


@pytest.fixture(scope="module")
def pdf_metrics() -> dict[int, tuple[float, float]]:
    return _probe_line_metrics()


class _LineCursor:
    """顺序排版游标：gap_before 为上一行 y1 到本行 y0 的净距离（pt）。"""

    def __init__(self, metrics: dict[int, tuple[float, float]]):
        self._metrics = metrics
        self.lines: list[tuple[str, float, float]] = []
        self._prev_y1 = 0.0

    def start(self, first_y0: float) -> None:
        self._prev_y1 = first_y0

    def add(self, text: str, size: int, gap_before: float) -> None:
        y0 = self._prev_y1 + gap_before
        self.lines.append((text, float(size), y0))
        self._prev_y1 = y0 + self._metrics[size][1]


def _compliant_pdf_pages(
    metrics: dict[int, tuple[float, float]],
    *,
    page1_pn_size: float | None = 14.0,
    page2_pn_size: float | None = 14.0,
) -> list[list[tuple[str, float, float]]]:
    """合规 2 页 PDF 的行布局（通知规则规格）。

    间距设计（与 _estimate_spacing_from_positions 的测量语义对齐）：
    - 段内行距统一 28.95pt（top-to-top），段内行间隙 g16 = 28.95 - 16pt行高；
    - 段后间距 = 行间净距 - g16：title/logo 段后 6pt 由 gap g16+6 产生，
      其余段落 gap 恒为 g16 → 段后 0pt；段前恒 None（PDF 语义）、首行缩进
      不测量（None 跳过）；页边距由 source="pdf" 跳过检查。
    """
    g16 = LINE_PITCH_PT - metrics[16][1]

    page1 = _LineCursor(metrics)
    page1.start(_TOP_PT)
    page1.add(_LOGO, 22, 0.0)  # header.issuing_logo（段后规范 6pt）
    page1.add(_NUMBER, 16, g16 + 6.0)  # header.issuing_number
    page1.add(_TITLE, 22, 40.0)  # body.title（版头/主体分界大空隙）
    page1.add(_ADDRESSEE, 16, g16 + 6.0)  # body.addressee（title 段后规范 6pt）
    for i in range(12):  # 正文 12 行，统一 28.95pt 节距，把成文日期推过 0.7 页高
        page1.add(f"正文内容第{i + 1}行，用于测试固定行距排版。", 16, g16)
    page1.add(_SIGNATURE, 16, g16)
    page1.add(_ISSUE_DATE, 16, g16)
    page1.add(_OFFICE_LINE, 14, 60.0)  # 主体/版记分界大空隙
    if page1_pn_size is not None:
        page1.add("1", int(page1_pn_size), g16)

    page2 = _LineCursor(metrics)
    page2.start(_TOP_PT)
    for i in range(4):
        page2.add(f"后续正文内容第{i + 1}行，保持固定行距排版。", 16, g16)
    page2.add(_SIGNATURE, 16, g16)
    page2.add(_ISSUE_DATE, 16, g16)
    page2.add(_OFFICE_LINE, 14, 60.0)
    if page2_pn_size is not None:
        page2.add("2", int(page2_pn_size), g16)

    return [page1.lines, page2.lines]


def _write_pdf(
    path: Path,
    pages_lines: list[list[tuple[str, float, float]]],
    metrics: dict[int, tuple[float, float]],
) -> None:
    doc = fitz.open()
    for lines in pages_lines:
        page = doc.new_page(width=_A4_W, height=_A4_H)
        for text, size, y0 in lines:
            ascent, _ = metrics[int(size)]
            page.insert_text((_LEFT_PT, y0 + ascent), text, fontname="china-s", fontsize=size)
    doc.save(str(path))
    doc.close()


class TestCompliantPdfE2E:
    """合规 PDF：零间距误报、零页边距误报、零页码误报。"""

    def test_no_spacing_margin_pagenumber_false_positives(
        self, tmp_path: Path, pdf_metrics: dict[int, tuple[float, float]]
    ) -> None:
        path = tmp_path / "compliant.pdf"
        _write_pdf(path, _compliant_pdf_pages(pdf_metrics), pdf_metrics)
        document, result = _audit(path)

        assert document.source == "pdf"
        assert document.total_page_num == 2
        assert _errors(result, "间距/缩进不匹配") == []
        assert _errors(result, "页边距不匹配") == []
        # 页码的存在性/页序语义零误报；字体类报错不在断言范围（测试环境
        # 无「仿宋」可内嵌，PDF 字体只能是 CJK 回退字体）。
        assert [
            e
            for e in result["errors"]
            if e["block_name"] == "页码" and e["error_type"] != "字体不匹配"
        ] == []
        # PDF 来源页边距检查应整体跳过并记入 unchecked
        assert any(u.get("check") == "margin" for u in result["unchecked"])


class TestPdfViolationE2E:
    """PDF 违规注入：按修后语义检出。"""

    def test_first_page_page_number_missing_detected(
        self, tmp_path: Path, pdf_metrics: dict[int, tuple[float, float]]
    ) -> None:
        """首页页码缺失报错（page_no=1）；第 2 页页码存在不误报缺失。"""
        path = tmp_path / "missing_pn.pdf"
        _write_pdf(path, _compliant_pdf_pages(pdf_metrics, page1_pn_size=None), pdf_metrics)
        _, result = _audit(path)
        errs = [e for e in _errors(result, "缺失必要元素") if e["block_name"] == "页码"]
        assert len(errs) == 1
        assert errs[0]["page_no"] == 1

    def test_later_page_page_number_wrong_size_detected(
        self, tmp_path: Path, pdf_metrics: dict[int, tuple[float, float]]
    ) -> None:
        """第 2 页页码字号 16pt（规范 14pt）→ 字体不匹配，page_no=2。"""
        path = tmp_path / "wrong_pn_size.pdf"
        _write_pdf(path, _compliant_pdf_pages(pdf_metrics, page2_pn_size=16.0), pdf_metrics)
        _, result = _audit(path)
        errs = [
            e
            for e in _errors(result, "字体不匹配")
            if e["block_name"] == "页码" and e["page_no"] == 2 and "字号应为 14" in e["details"]
        ]
        assert len(errs) == 1


# ---------------------------------------------------------------------------
# DOCX 生成
# ---------------------------------------------------------------------------


def _write_docx(path: Path, *, bad_space_before: float | None = None) -> None:
    """生成合规 DOCX（通知规则规格）；bad_space_before 注入第 2 个正文段的段前距。"""
    doc = DocxDocument()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(37), Mm(35)
    section.left_margin, section.right_margin = Mm(28), Mm(26)

    def add(
        text: str,
        font: str,
        size: float,
        *,
        bold: bool = False,
        align: str | None = None,
        before: float = 0.0,
        after: float = 0.0,
        line: float = LINE_PITCH_PT,
        indent: float = 0.0,
    ) -> None:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = font
        run.font.size = Pt(size)
        run.bold = bold
        pf = p.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = Pt(line)
        pf.first_line_indent = Pt(indent)
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    add(_LOGO, "方正小标宋简体", 22, bold=True, align="center", before=12, after=6, line=33.0)
    add(_NUMBER, "仿宋", 16, before=3, after=3, line=16.0)
    add(_TITLE, "方正小标宋简体", 22, align="center", before=12, after=6, line=33.0)
    add(_ADDRESSEE, "仿宋", 16)
    for i in range(3):
        before = bad_space_before if (bad_space_before is not None and i == 1) else 0.0
        add(
            f"正文内容第{i + 1}段，用于测试固定行距与首行缩进。",
            "仿宋",
            16,
            before=before,
            indent=32.0,
        )
    add(_SIGNATURE, "仿宋", 16, align="right", before=12)
    add(_ISSUE_DATE, "仿宋", 16, align="right", before=3)
    add(_OFFICE_LINE, "仿宋", 14, before=3)

    # 页码在 Word 节页脚中（docx_parser 会提取并标记为 _word_footer）
    footer_p = section.footer.paragraphs[0]
    run = footer_p.add_run("1")
    run.font.name = "仿宋"
    run.font.size = Pt(14)
    run.bold = False
    pf = footer_p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(0)
    pf.first_line_indent = Pt(0)

    doc.save(str(path))


class TestCompliantDocxE2E:
    """合规 DOCX：图形要素缺失报错属预期，其余零错误。"""

    def test_only_graphic_blocks_reported_missing(self, tmp_path: Path) -> None:
        path = tmp_path / "compliant.docx"
        _write_docx(path)
        document, result = _audit(path)

        assert document.source == "docx"
        # 红色分隔线/印章/黑色反线是图形而非文本行，规则引擎无法分类，
        # 缺失必要元素报错属预期；除此之外不允许有任何错误。
        assert sorted(e["block_name"] for e in result["errors"]) == [
            "印章",
            "红色分隔线",
            "黑色反线",
        ]
        assert all(e["error_type"] == "缺失必要元素" for e in result["errors"])


class TestDocxViolationE2E:
    def test_space_before_violation_detected(self, tmp_path: Path) -> None:
        """正文段段前距 20pt（规范 0pt）→ 检出间距/缩进不匹配。"""
        path = tmp_path / "bad_spacing.docx"
        _write_docx(path, bad_space_before=20.0)
        _, result = _audit(path)
        errs = _errors(result, "间距/缩进不匹配")
        assert len(errs) == 1
        assert "段前间距应为 0pt" in errs[0]["details"]
        assert "实际为 20pt" in errs[0]["details"]
