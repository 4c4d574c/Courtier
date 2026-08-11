"""Tests for DOCX body-order table extraction, merged-cell dedup,
and Word header/footer extraction."""

from __future__ import annotations

from docparse.parsers.base import ParserConfig
from docparse.parsers.docx_parser import (
    _extract_all_lines,
    _extract_header_footer_lines,
    _FontResolver,
)
from docparse.parsers.registry import parse
from docparse.parsers.scanned.structure import collect_all_paragraphs
from docx import Document as DocxDocument
from docx.shared import Pt
from lxml import etree

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _resolver(doc) -> _FontResolver:
    return _FontResolver(doc, {})


def _texts(lines: list[dict]) -> list[str]:
    return [line["text"] for line in lines if line.get("text")]


def _add_page_break(para) -> None:
    br = etree.SubElement(para.runs[0]._element, f"{{{_W_NS}}}br")
    br.set(f"{{{_W_NS}}}type", "page")


class TestTableExtraction:
    def test_table_text_appears_in_body_order(self):
        doc = DocxDocument()
        doc.add_paragraph("第一段")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "表格内容"
        doc.add_paragraph("第二段")

        texts = _texts(_extract_all_lines(doc, _resolver(doc)))

        assert texts.index("第一段") < texts.index("表格内容") < texts.index("第二段")

    def test_horizontally_merged_cell_extracted_once(self):
        doc = DocxDocument()
        table = doc.add_table(rows=2, cols=2)
        merged = table.cell(0, 0).merge(table.cell(0, 1))
        merged.text = "合并单元格"
        table.cell(1, 0).text = "甲"
        table.cell(1, 1).text = "乙"

        texts = _texts(_extract_all_lines(doc, _resolver(doc)))

        assert texts.count("合并单元格") == 1
        assert "甲" in texts
        assert "乙" in texts

    def test_vertically_merged_cell_extracted_once(self):
        doc = DocxDocument()
        table = doc.add_table(rows=2, cols=2)
        merged = table.cell(0, 0).merge(table.cell(1, 0))
        merged.text = "纵向合并"
        table.cell(0, 1).text = "丙"
        table.cell(1, 1).text = "丁"

        texts = _texts(_extract_all_lines(doc, _resolver(doc)))

        assert texts.count("纵向合并") == 1
        assert "丙" in texts
        assert "丁" in texts


class TestHeaderFooterExtraction:
    def test_header_footer_lines_tagged(self):
        doc = DocxDocument()
        doc.add_paragraph("正文内容")
        section = doc.sections[0]
        section.header.paragraphs[0].text = "秘密★1年"
        section.footer.paragraphs[0].text = "— 1 —"

        header_lines, footer_lines = _extract_header_footer_lines(doc, _resolver(doc))

        assert [line["text"] for line in header_lines] == ["秘密★1年"]
        assert [line["text"] for line in footer_lines] == ["— 1 —"]
        assert all(line.get("_word_header") for line in header_lines)
        assert all(line.get("_word_footer") for line in footer_lines)

    def test_parse_classifies_header_footer(self, tmp_path):
        doc = DocxDocument()
        doc.add_paragraph("现将有关事项通知如下。")
        section = doc.sections[0]
        section.header.paragraphs[0].text = "秘密★1年"
        section.footer.paragraphs[0].text = "— 1 —"
        path = tmp_path / "doc.docx"
        doc.save(str(path))

        result = parse(str(path), ParserConfig())

        page_content = result.pages[0].page_content
        classification = page_content.header.classification_duration
        assert classification is not None
        assert classification.elements[0].font.text == "秘密★1年"
        assert page_content.footer.page_number.elements[0].font.text == "— 1 —"

    def test_header_footer_deduped_across_pages(self, tmp_path):
        """Word headers/footers attach to the first page group only."""
        doc = DocxDocument()
        para1 = doc.add_paragraph("第一页内容")
        _add_page_break(para1)
        doc.add_paragraph("第二页内容")
        section = doc.sections[0]
        section.footer.paragraphs[0].text = "— 1 —"
        path = tmp_path / "multi.docx"
        doc.save(str(path))

        result = parse(str(path), ParserConfig())

        assert result.total_page_num == 2
        first_footer = result.pages[0].page_content.footer
        second_footer = result.pages[1].page_content.footer
        assert first_footer.page_number is not None
        assert first_footer.page_number.elements[0].font.text == "— 1 —"
        assert second_footer.page_number is None

    def test_pages_carry_no_raw_or_save_path(self, tmp_path):
        """Page.raw / Page.save_path were removed from the model."""
        from docmodels import Page

        doc = DocxDocument()
        doc.add_paragraph("正文内容")
        path = tmp_path / "plain.docx"
        doc.save(str(path))

        result = parse(str(path), ParserConfig())

        assert "raw" not in Page.model_fields
        assert "save_path" not in Page.model_fields
        assert result.pages


class TestParagraphSpacingExtraction:
    """DOCX paragraph spacing/indent is read directly from python-docx."""

    def _make_spaced_doc(self, line_spacing) -> DocxDocument:
        doc = DocxDocument()
        para = doc.add_paragraph()
        run = para.add_run("现将有关事项通知如下，请各部门结合实际抓好贯彻落实。")
        run.font.size = Pt(16)
        pf = para.paragraph_format
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)
        pf.line_spacing = line_spacing
        pf.first_line_indent = Pt(32)
        pf.left_indent = Pt(10)
        pf.right_indent = Pt(5)
        return doc

    def test_spacing_keys_on_extracted_lines(self):
        doc = self._make_spaced_doc(line_spacing=1.5)

        lines = _extract_all_lines(doc, _resolver(doc))

        assert len(lines) == 1
        line = lines[0]
        assert line["space_before"] == 12.0
        assert line["space_after"] == 6.0
        # 1.5x multiple × 16pt max font size → 24pt
        assert line["line_spacing"] == 24.0
        assert line["first_indent"] == 32.0
        assert line["left_indent"] == 10.0
        assert line["right_indent"] == 5.0

    def test_absolute_line_spacing_used_directly(self):
        doc = self._make_spaced_doc(line_spacing=Pt(28))

        lines = _extract_all_lines(doc, _resolver(doc))

        assert lines[0]["line_spacing"] == 28.0

    def test_unset_spacing_stays_none(self):
        doc = DocxDocument()
        doc.add_paragraph("现将有关事项通知如下，请各部门结合实际抓好贯彻落实。")

        lines = _extract_all_lines(doc, _resolver(doc))

        assert len(lines) == 1
        for key in (
            "space_before",
            "space_after",
            "line_spacing",
            "first_indent",
            "left_indent",
            "right_indent",
        ):
            assert lines[0][key] is None

    def test_parse_populates_paragraph_spacing(self, tmp_path):
        doc = self._make_spaced_doc(line_spacing=1.5)
        path = tmp_path / "spaced.docx"
        doc.save(str(path))

        result = parse(str(path), ParserConfig())

        paras = [p for page in result.pages for p in collect_all_paragraphs(page.page_content)]
        target = [p for p in paras if p.space_before == 12.0]
        assert len(target) == 1
        assert target[0].space_after == 6.0
        assert target[0].line_spacing == 24.0
        assert target[0].first_indent == 32.0
        assert target[0].left_indent == 10.0
        assert target[0].right_indent == 5.0

    def test_parse_unset_spacing_stays_none(self, tmp_path):
        doc = DocxDocument()
        doc.add_paragraph("现将有关事项通知如下，请各部门结合实际抓好贯彻落实。")
        path = tmp_path / "plain.docx"
        doc.save(str(path))

        result = parse(str(path), ParserConfig())

        paras = [p for page in result.pages for p in collect_all_paragraphs(page.page_content)]
        assert paras
        assert all(p.space_before is None for p in paras)
        assert all(p.line_spacing is None for p in paras)


class TestDocxParagraphsNotMerged:
    """DOCX 不合并：每个 w:p 本来就是真实段落，解析后各自间距完整。"""

    def test_three_word_paragraphs_keep_own_spacing(self, tmp_path):
        """3 个 Word 段、中间段 20pt space_before → 解析后 3 段各自间距完整。"""
        doc = DocxDocument()
        doc.add_paragraph("为贯彻落实上级决策部署，扎实推进各项工作有序开展。")
        para2 = doc.add_paragraph("各部门单位要高度重视，切实履行主体责任确保到位。")
        pf2 = para2.paragraph_format
        pf2.space_before = Pt(20)
        pf2.space_after = Pt(8)
        doc.add_paragraph("请结合实际认真抓好落实，确保各项任务按时完成。")
        path = tmp_path / "three_paras.docx"
        doc.save(str(path))

        result = parse(str(path), ParserConfig())

        main_text = result.pages[0].page_content.body.main_text
        assert len(main_text) == 3
        # 每个 w:p 独立成段（各一行），不被视觉合并吞掉
        assert all(len(p.elements) == 1 for p in main_text)
        assert main_text[0].space_before is None
        assert main_text[1].space_before == 20.0
        assert main_text[1].space_after == 8.0
        assert main_text[2].space_before is None

    def test_parse_source_is_docx(self, tmp_path):
        """registry.parse 回填 Document.source='docx'。"""
        doc = DocxDocument()
        doc.add_paragraph("正文内容")
        path = tmp_path / "plain.docx"
        doc.save(str(path))

        result = parse(str(path), ParserConfig())

        assert result.source == "docx"
        assert result.schema_version == "1.0"


class TestUnclassifiedWarningsSurfaced:
    """规则引擎未归类行的摘要应进入 Document.warnings（修前只进日志）。"""

    def test_docx_unclassified_lines_in_document_warnings(self, tmp_path, monkeypatch):
        from types import SimpleNamespace

        import docparse.parsers.rules as rules_mod
        from docparse.parsers.docx_parser import DocxParser
        from docparse.parsers.rules import ClassifiedLine

        src = DocxDocument()
        src.add_paragraph("一行用于测试警告透传的普通正文内容。")
        path = tmp_path / "unclassified.docx"
        src.save(str(path))

        class _StubEngine:
            """全部行强制 unclassified，隔离规则引擎启发式的不确定性。"""

            def classify_lines(self, lines, has_position):
                return SimpleNamespace(
                    lines=[
                        ClassifiedLine(
                            line_no=i,
                            text=line.get("text", ""),
                            field="unclassified",
                            confidence=0.0,
                        )
                        for i, line in enumerate(lines)
                    ]
                )

        # docx_parser 在函数内 import StructureRuleEngine，补丁打在源模块上。
        monkeypatch.setattr(rules_mod, "StructureRuleEngine", _StubEngine)

        doc = DocxParser().parse(str(path), ParserConfig())

        assert any(
            w.startswith("第 1 页：") and "未能自动归类" in w for w in doc.warnings
        ), f"未归类摘要应进入 Document.warnings：{doc.warnings}"
