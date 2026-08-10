"""Tests for parser registry and file type detection."""

import pytest
from docparse.parsers.base import ParserConfig
from docparse.parsers.registry import (
    _file_extension,
    get_parser,
)
from lxml import etree

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class TestParserConfig:
    """Test ParserConfig creation and defaults."""

    def test_default_config(self):
        config = ParserConfig()
        assert config.llm_base_url == ""
        assert config.llm_api_key == ""
        assert config.llm_model == ""
        assert config.ocr_lang == "ch"
        assert config.ocr_api_url == ""

    def test_custom_config(self):
        config = ParserConfig(
            llm_base_url="http://localhost:8000/v1",
            llm_api_key="test-key",
            llm_model="gpt-4",
        )
        assert config.llm_base_url == "http://localhost:8000/v1"
        assert config.llm_api_key == "test-key"
        assert config.llm_model == "gpt-4"


class TestFileTypeDetection:
    """Test file extension detection."""

    def test_pdf_extension(self):
        assert _file_extension("test.pdf") == ".pdf"
        assert _file_extension("test.PDF") == ".pdf"
        assert _file_extension("/path/to/my file.pdf") == ".pdf"

    def test_docx_extension(self):
        assert _file_extension("document.docx") == ".docx"
        assert _file_extension("DOCUMENT.DOCX") == ".docx"

    def test_image_extensions(self):
        assert _file_extension("scan.png") == ".png"
        assert _file_extension("photo.jpg") == ".jpg"
        assert _file_extension("scan.jpeg") == ".jpeg"
        assert _file_extension("doc.tiff") == ".tiff"
        assert _file_extension("doc.bmp") == ".bmp"

    def test_unknown_extension(self):
        assert _file_extension("file.xyz") == ".xyz"


class TestGetParser:
    """Test parser selection based on file type."""

    def test_docx_parser_for_docx(self, tmp_path):
        from docparse.parsers.docx_parser import DocxParser

        docx_file = tmp_path / "test.docx"
        docx_file.write_text("dummy")
        parser = get_parser(str(docx_file))
        assert isinstance(parser, DocxParser)

    def test_scanned_parser_for_image(self, tmp_path):
        from docparse.parsers.scanned import ScannedParser

        png_file = tmp_path / "test.png"
        png_file.write_bytes(b"dummy")
        parser = get_parser(str(png_file))
        assert isinstance(parser, ScannedParser)

    def test_unsupported_file_type(self, tmp_path):
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("dummy")
        with pytest.raises(ValueError, match="Unsupported file type"):
            get_parser(str(txt_file))

    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            from docparse.parsers.registry import parse

            parse("/nonexistent/file.pdf")


class TestParseSourceBackfill:
    """registry.parse 在分派后回填 Document.source。"""

    def test_image_source_is_scanned(self, tmp_path, monkeypatch):
        import docparse.parsers.scanned as scanned_mod
        from docmodels import Document
        from docparse.parsers.registry import parse

        class _FakeScannedParser:
            def parse(self, file_path, config=None):
                return Document(doc_id="fake-img")

        monkeypatch.setattr(scanned_mod, "ScannedParser", _FakeScannedParser)
        img = tmp_path / "scan.png"
        img.write_bytes(b"dummy")

        result = parse(str(img), ParserConfig())

        assert result.doc_id == "fake-img"
        assert result.source == "scanned"


class TestDocxPageBreakDetection:
    """Test page break detection in DOCX paragraphs."""

    def test_has_page_break_with_explicit_break(self, tmp_path):
        from docparse.parsers.docx_parser import _has_page_break
        from docx import Document as DocxDocument

        doc = DocxDocument()
        para = doc.add_paragraph("Hello")
        run = para.add_run()
        run._element.append(
            etree.SubElement(run._element, f"{{{_W_NS}}}br", attrib={f"{{{_W_NS}}}type": "page"})
        )
        assert _has_page_break(para) is True

    def test_has_page_break_without_break(self, tmp_path):
        from docparse.parsers.docx_parser import _has_page_break
        from docx import Document as DocxDocument

        doc = DocxDocument()
        para = doc.add_paragraph("Hello")
        assert _has_page_break(para) is False

    def test_extract_lines_with_page_breaks(self, tmp_path):
        from docparse.parsers.docx_parser import _extract_all_lines, _FontResolver
        from docx import Document as DocxDocument

        doc = DocxDocument()
        para1 = doc.add_paragraph("First page")
        run1 = para1.add_run()
        br_elem = etree.SubElement(run1._element, f"{{{_W_NS}}}br")
        br_elem.set(f"{{{_W_NS}}}type", "page")
        doc.add_paragraph("Second page")

        theme_fonts = {}
        font_resolver = _FontResolver(doc, theme_fonts)
        lines = _extract_all_lines(doc, font_resolver)

        page_break_indices = [i for i, line in enumerate(lines) if line.get("_page_break")]
        assert len(page_break_indices) == 1
        assert lines[page_break_indices[0] - 1]["text"] == "First page"


class TestSplitLinesByPageBreaks:
    """Test splitting lines by page break markers."""

    def test_no_page_breaks(self):
        from docparse.parsers.docx_parser import _split_lines_by_page_breaks

        lines = [
            {"text": "Line 1", "line_no": 0},
            {"text": "Line 2", "line_no": 1},
        ]
        pages = _split_lines_by_page_breaks(lines)
        assert len(pages) == 1
        assert len(pages[0]) == 2

    def test_one_page_break(self):
        from docparse.parsers.docx_parser import _split_lines_by_page_breaks

        lines = [
            {"text": "Line 1", "line_no": 0},
            {"_page_break": True, "line_no": 1},
            {"text": "Line 2", "line_no": 2},
        ]
        pages = _split_lines_by_page_breaks(lines)
        assert len(pages) == 2
        assert len(pages[0]) == 1
        assert pages[0][0]["text"] == "Line 1"
        assert len(pages[1]) == 1
        assert pages[1][0]["text"] == "Line 2"

    def test_empty_lines(self):
        from docparse.parsers.docx_parser import _split_lines_by_page_breaks

        pages = _split_lines_by_page_breaks([])
        assert len(pages) == 1
        assert pages[0] == []

    def test_consecutive_page_breaks(self):
        from docparse.parsers.docx_parser import _split_lines_by_page_breaks

        lines = [
            {"text": "Line 1", "line_no": 0},
            {"_page_break": True, "line_no": 1},
            {"_page_break": True, "line_no": 2},
            {"text": "Line 2", "line_no": 3},
        ]
        pages = _split_lines_by_page_breaks(lines)
        assert len(pages) == 3
        assert len(pages[0]) == 1
        assert pages[1] == []
        assert len(pages[2]) == 1


class TestDocxMultiPage:
    """Test DOCX parser produces multi-page output."""

    def test_single_page_docx(self, tmp_path):
        from docparse.parsers.base import ParserConfig
        from docparse.parsers.registry import parse
        from docx import Document as DocxDocument

        doc = DocxDocument()
        doc.add_paragraph("Hello World")
        docx_path = tmp_path / "single.docx"
        doc.save(str(docx_path))

        config = ParserConfig()
        result = parse(str(docx_path), config)
        assert result.total_page_num >= 1
        assert len(result.pages) >= 1
        assert result.pages[0].page_no == 0

    def test_page_break_produces_multiple_pages(self, tmp_path):
        from docparse.parsers.base import ParserConfig
        from docparse.parsers.registry import parse
        from docx import Document as DocxDocument

        doc = DocxDocument()
        para1 = doc.add_paragraph("First page content")
        run1 = para1.runs[0] if para1.runs else para1.add_run("First page content")
        br = etree.SubElement(run1._element, f"{{{_W_NS}}}br")
        br.set(f"{{{_W_NS}}}type", "page")
        doc.add_paragraph("Second page content")
        docx_path = tmp_path / "multipage.docx"
        doc.save(str(docx_path))

        config = ParserConfig()
        result = parse(str(docx_path), config)
        assert result.total_page_num == 2
        assert len(result.pages) == 2
        assert result.pages[0].page_no == 0
        assert result.pages[1].page_no == 1
